"""Renderer: checked deck -> editable .pptx with native (editable) charts.

Chart values are looked up from the fact table by ID, never taken from model text.
Anything the fact-checker could not verify is marked in red on the slide.
"""
from io import BytesIO

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.oxml.ns import qn
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from app.deck import Bullet, Deck, Slide
from app.factcheck import CheckReport
from app.facts import FactTable

PINK = RGBColor(0xFF, 0x2E, 0x93)
BLACK = RGBColor(0x0B, 0x0B, 0x0D)
GRAY = RGBColor(0x6B, 0x6B, 0x75)
LIGHT = RGBColor(0xD4, 0xD4, 0xD8)
RED = RGBColor(0xDC, 0x26, 0x26)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
W, H = Inches(13.333), Inches(7.5)


def _text(slide, x, y, w, h, text, size=14, bold=False, color=BLACK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, color
    run.font.name = "Calibri"
    return tf


def _bar(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(1, x, y, w, h)  # 1 = rectangle
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def _bullets(slide, x, y, w, h, bullets: list[Bullet], bad: set[str], prefix: str, size=16):
    tf = slide.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets, 1):
        p = tf.paragraphs[0] if i == 1 else tf.add_paragraph()
        p.space_after = Pt(10)
        dot = p.add_run()
        dot.text = "■  "
        dot.font.size, dot.font.color.rgb = Pt(size - 6), PINK
        run = p.add_run()
        run.text = b.text
        run.font.size, run.font.color.rgb, run.font.name = Pt(size), BLACK, "Calibri"
        if f"{prefix} bullet {i}" in bad:
            flag = p.add_run()
            flag.text = "  [UNVERIFIED NUMBER]"
            flag.font.size, flag.font.bold, flag.font.color.rgb = Pt(size - 4), True, RED


def _number_format(unit: str, values: list[float]) -> str:
    if unit == "pct":
        return '+0.0"%";-0.0"%";0.0"%"' if any(v < 0 for v in values) else '0.0"%"'
    if unit == "pp":
        return '+0.0" pp";-0.0" pp"'
    if unit == "usd":
        big = max((abs(v) for v in values), default=0)
        body = '$#,##0.00,,"M"' if big >= 1e6 else ('$#,##0,"K"' if big >= 1e3 else "$#,##0")
        return f"+{body};-{body}" if any(v < 0 for v in values) else body
    return "#,##0"


def _chart(slide, s: Slide, facts: FactTable, x, y, w, h) -> None:
    points = [(p.label, facts.get(p.fact_id)) for p in s.chart.points if facts.get(p.fact_id)]
    if not points:
        return
    data = CategoryChartData()
    data.categories = [label for label, _ in points]
    values = [f.value for _, f in points]
    data.add_series(s.chart.title, values)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if s.chart.type == "bar" else XL_CHART_TYPE.COLUMN_CLUSTERED
    chart = slide.shapes.add_chart(kind, x, y, w, h, data).chart
    chart.has_legend = False
    chart.has_title = True
    chart.chart_title.text_frame.text = s.chart.title
    chart.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
    chart.chart_title.text_frame.paragraphs[0].runs[0].font.color.rgb = GRAY
    plot = chart.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = _number_format(points[0][1].unit, values)
    labels.number_format_is_linked = False
    labels.font.size = Pt(12)
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    series = plot.series[0]
    series.invert_if_negative = False
    for i, v in enumerate(values):
        pt = series.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = PINK if v >= 0 else BLACK
        # PowerPoint renders negative bars hollow unless each data point opts out explicitly.
        dpt = pt.format._element  # the <c:dPt> for this point
        if dpt.find(qn("c:invertIfNegative")) is None:
            inv = dpt.makeelement(qn("c:invertIfNegative"), {"val": "0"})
            dpt.insert(1, inv)  # schema order: idx, invertIfNegative, marker, bubble3D, explosion, spPr
    value_axis = chart.value_axis
    value_axis.visible = False
    value_axis.has_major_gridlines = False
    chart.category_axis.tick_labels.font.size = Pt(12)
    chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW  # labels clear of negative bars
    if s.chart.type == "bar":
        chart.category_axis.reverse_order = True  # biggest first, top to bottom


def _frame(prs, number: int, total: int, footer: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _bar(slide, 0, 0, W, Inches(0.12), PINK)
    _text(slide, Inches(0.6), H - Inches(0.5), Inches(9), Inches(0.3), footer, size=9, color=GRAY)
    _text(slide, W - Inches(1.6), H - Inches(0.5), Inches(1.0), Inches(0.3), f"{number} / {total}",
          size=9, color=GRAY, align=PP_ALIGN.RIGHT)
    return slide


def render_pptx(deck: Deck, facts: FactTable, check: CheckReport) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    bad = {v["where"] for v in check.violations}
    total = len(deck.slides) + 2
    status = ("All numbers fact-checked against the source data"
              if check.passed else f"{len(check.violations)} number(s) could not be verified (marked in red)")
    footer = f"Source: client data files. {status}. Generated by Deck Autopilot."

    # Title slide: black, pink accent.
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BLACK
    _bar(s, Inches(0.8), Inches(2.6), Inches(1.2), Inches(0.1), PINK)
    _text(s, Inches(0.8), Inches(2.8), Inches(11.5), Inches(1.6), deck.deck_title, size=40, bold=True, color=WHITE)
    _text(s, Inches(0.8), Inches(4.3), Inches(11.5), Inches(0.6), deck.subtitle, size=20, color=PINK)
    _text(s, Inches(0.8), H - Inches(0.9), Inches(11.5), Inches(0.4), status, size=12, color=LIGHT)

    # Executive summary.
    s = _frame(prs, 2, total, footer)
    _text(s, Inches(0.6), Inches(0.45), Inches(12), Inches(0.9), "Executive summary", size=30, bold=True)
    _bullets(s, Inches(0.6), Inches(1.5), Inches(12), Inches(5.3), deck.executive_summary, bad,
             "executive summary", size=20)

    for i, slide in enumerate(deck.slides, 1):
        s = _frame(prs, i + 2, total, footer)
        title_color = RED if f"slide {i} title" in bad else BLACK
        _text(s, Inches(0.6), Inches(0.4), Inches(12.1), Inches(1.2), slide.title, size=26, bold=True,
              color=title_color)
        has_chart = slide.chart is not None and any(facts.get(p.fact_id) for p in slide.chart.points)
        text_w = Inches(5.6) if has_chart else Inches(12)
        _bullets(s, Inches(0.6), Inches(1.8), text_w, Inches(5.0), slide.bullets, bad, f"slide {i}")
        if has_chart:
            _chart(s, slide, facts, Inches(6.6), Inches(1.7), Inches(6.2), Inches(5.0))
        cited = sorted({f for b in slide.bullets for f in b.fact_ids} | set(slide.title_fact_ids),
                       key=lambda x: int(x[1:]) if x[1:].isdigit() else 0)
        trail = "\n".join(f"{fid}: {facts.get(fid).display} = {facts.get(fid).label}"
                          for fid in cited if facts.get(fid))
        s.notes_slide.notes_text_frame.text = f"{slide.speaker_notes}\n\nFacts on this slide:\n{trail}"

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def chart_values(deck: Deck, facts: FactTable) -> list[dict | None]:
    """Resolve chart fact IDs to values for the web preview (same source as the .pptx)."""
    out = []
    for s in deck.slides:
        if not s.chart:
            out.append(None)
            continue
        pts = [{"label": p.label, "fact_id": p.fact_id, "value": facts.get(p.fact_id).value,
                "display": facts.get(p.fact_id).display, "unit": facts.get(p.fact_id).unit}
               for p in s.chart.points if facts.get(p.fact_id)]
        out.append({"type": s.chart.type, "title": s.chart.title, "points": pts} if pts else None)
    return out

