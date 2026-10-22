"""Offline tests for the .pptx renderer, using a saved real run (tests/fixtures/sample_run.json)."""
import json
from io import BytesIO
from pathlib import Path

from pptx import Presentation

from app.deck import Deck
from app.factcheck import CheckReport
from app.facts import Fact, FactTable
from app.render_pptx import chart_values, render_pptx

FIXTURE = Path(__file__).parent / "fixtures" / "sample_run.json"


def _load():
    r = json.loads(FIXTURE.read_text(encoding="utf-8"))
    facts = FactTable()
    for f in r["facts"]:
        facts.facts[f["id"]] = Fact(f["id"], f["label"], f["value"], f["unit"], f["source"])
    return Deck.model_validate(r["deck"]), facts


def _texts(prs):
    return [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]


def test_renders_title_summary_and_content_slides():
    deck, facts = _load()
    prs = Presentation(BytesIO(render_pptx(deck, facts, CheckReport())))
    assert len(prs.slides) == len(deck.slides) + 2
    texts = _texts(prs)
    assert deck.deck_title in texts[0] and "Executive summary" in texts[1]
    assert "UNVERIFIED" not in " ".join(texts)


def test_chart_values_come_from_fact_table_not_model_text():
    deck, facts = _load()
    charted = [(s, c) for s, c in zip(deck.slides, chart_values(deck, facts)) if c]
    assert charted
    slide, chart = charted[0]
    for p, point in zip(slide.chart.points, chart["points"]):
        assert point["value"] == facts.get(p.fact_id).value
    prs = Presentation(BytesIO(render_pptx(deck, facts, CheckReport())))
    native = [sh.chart for s in prs.slides for sh in s.shapes if sh.has_chart]
    assert len(native) == len(charted)
    assert list(native[0].plots[0].series[0].values) == [p["value"] for p in chart["points"]]


def test_unverified_numbers_are_flagged_on_the_slide():
    deck, facts = _load()
    report = CheckReport(violations=[{"where": "slide 1 bullet 1", "text": "x", "number": "93.7%",
                                      "reason": "no fact in the table has this value"}])
    prs = Presentation(BytesIO(render_pptx(deck, facts, report)))
    texts = _texts(prs)
    assert "[UNVERIFIED NUMBER]" in texts[2]
    assert "could not be verified" in texts[0]


def test_speaker_notes_carry_the_fact_trail():
    deck, facts = _load()
    prs = Presentation(BytesIO(render_pptx(deck, facts, CheckReport())))
    notes = prs.slides[2].notes_slide.notes_text_frame.text
    assert "Facts on this slide:" in notes and "F" in notes
